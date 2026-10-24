# -*- coding: utf-8 -*-
"""
Happy2Help — judges' pitch deck builder.
One slide spec → (1) editable PPTX via python-pptx, (2) an HTML mirror used only to render PNG previews + PDF.
All images are the real brand assets / real app screenshots from this repository. No synthetic UI.
Run:  python presentation/build_pitch.py
"""
import os, sys, json, html
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "presentation")
W, H = 1920, 1080            # design pixels
EMU_PER_PX = 914400 * 13.333333 / W   # 13.333 in wide slide

NAVY, NAVY_DEEP, SKY, SKY_DEEP = "0F2D5B", "081C3A", "5B9CFF", "2D67B5"
SUN, MINT, CLOUD, WHITE, INK, MUTED, RED = "FFD34D", "A7E3C1", "F3F6FB", "FFFFFF", "173153", "8FA3BC", "D64550"
FONT = "Segoe UI"

def A(*p): return os.path.join(ROOT, *p)
BRAND = {
    "wordmark": A("web", "public", "brand", "wordmark.png"),
    "wordmark_navy": A("web", "public", "brand", "wordmark-on-navy.png"),
    "mascot": A("web", "public", "brand", "mascot.png"),
    "mark": A("web", "public", "brand", "mark.png"),
}
SHOT = lambda *p: A("artifacts", "screenshots", *p)

# ------------------------------------------------------------------ slide spec
# element kinds: bg, rect, text, img, ring, line
slides = []

def S(bg, notes, *els):
    slides.append({"bg": bg, "notes": notes, "els": list(els)})

def T(x, y, w, h, text, size, color=NAVY, bold=False, align="l", spacing=None, line=1.15, anchor="t", font=FONT):
    return {"k": "text", "x": x, "y": y, "w": w, "h": h, "text": text, "size": size, "color": color, "bold": bold, "align": align, "spacing": spacing, "line": line, "anchor": anchor, "font": font}
def I(path, x, y, w, h, radius=0, fit="cover"):
    return {"k": "img", "path": path, "x": x, "y": y, "w": w, "h": h, "radius": radius, "fit": fit}
def R(x, y, w, h, fill, radius=0, line=None, alpha=None):
    return {"k": "rect", "x": x, "y": y, "w": w, "h": h, "fill": fill, "radius": radius, "line": line, "alpha": alpha}
def RING(cx, cy, r, color, width=3):
    return {"k": "ring", "cx": cx, "cy": cy, "r": r, "color": color, "width": width}
def DOT(cx, cy, r, color):
    return {"k": "dot", "cx": cx, "cy": cy, "r": r, "color": color}

# 01 — THE PROBLEM (quiet, navy, one unanswered signal)
S(NAVY_DEEP,
  "Тусламж хэрэгтэй үед… (2 секунд түр зогс) …хүүхэд хэнд хандах вэ? Нэг хүүхэд тусламж хэрэгтэй гэдгээ мэдэж байна. Харин хэнд хэлэхээ, яаж хэлэхээ мэдэхгүй байна.",
  T(160, 330, 1200, 130, "Тусламж хэрэгтэй үед…", 84, WHITE, False),
  T(160, 470, 1300, 150, "Хүүхэд хэнд хандах вэ?", 96, WHITE, True),
  RING(1480, 560, 60, SKY_DEEP), RING(1480, 560, 120, SKY_DEEP, 2), RING(1480, 560, 190, SKY_DEEP, 1),
  DOT(1480, 560, 14, SKY),
  T(160, 940, 900, 40, "Хариугүй үлдсэн дохио.", 24, MUTED, False),
)

# 02 — THE REAL CHALLENGE
S(NAVY,
  "Мэдээлэх нь эхлэл. Харин тусламж үнэхээр хүрсэн үү? Гурван асуулт байна. Хэн сонсох вэ? Хэн хариулах вэ? Дараа нь юу болох вэ? Сурагч мэдээлнэ, зөв насанд хүрэгч хүлээж авах ёстой, сурагч хүлээж авсныг мэдэх ёстой. Өнөөдөр энэ гурав ихэвчлэн тасарсан байдаг.",
  T(160, 150, 1600, 220, "Мэдээлэх нь эхлэл.\nХарин тусламж үнэхээр хүрсэн үү?", 64, WHITE, True, line=1.15),
  R(160, 460, 500, 380, "123F73", radius=28), R(710, 460, 500, 380, "123F73", radius=28), R(1260, 460, 500, 380, "123F73", radius=28),
  T(200, 500, 420, 100, "Хэн сонсох вэ?", 40, WHITE, True),
  T(200, 620, 420, 180, "Сурагч айдастай, хэнд хэлэхээ мэдэхгүй.", 26, "C9D8EE"),
  T(750, 500, 420, 100, "Хэн хариулах вэ?", 40, WHITE, True),
  T(750, 620, 420, 180, "Зөв насанд хүрэгч мэдээллийг цаг алдалгүй авах ёстой.", 26, "C9D8EE"),
  T(1300, 500, 420, 100, "Дараа нь юу болох вэ?", 40, WHITE, True),
  T(1300, 620, 420, 180, "Сурагч хүлээж авсныг, хэн хариуцаж байгааг мэдэх ёстой.", 26, "C9D8EE"),
  T(160, 930, 1600, 50, "Эдгээр гурван холбоос тасрахад мэдээлэл замдаа алга болдог.", 26, MUTED),
)

# 03 — THAT'S WHY WE BUILT (signature moment, bright)
S(WHITE,
  "That's why we built… Happy2Help. (Нэрийг хэлээд 2 секунд түр зогс.) Хүүхэд бүр сонсогдох эрхтэй.",
  T(160, 250, 1100, 70, "THAT'S WHY WE BUILT…", 36, SKY_DEEP, True, spacing=6),
  I(BRAND["wordmark"], 160, 360, 900, 306, fit="contain"),
  I(BRAND["mascot"], 1250, 220, 520, 601, fit="contain"),
  T(160, 740, 1100, 100, "Хүүхэд бүр сонсогдох эрхтэй.", 56, NAVY, True),
  T(160, 860, 1000, 40, "Сургуулийн аюулгүй байдал, итгэлцлийн платформ · erxes plugin · нээлттэй эх", 22, MUTED),
)

# 04 — THE SOLUTION (real SOS phone + real teacher dashboard)
S(WHITE,
  "Нэг дохионоос бодит тусламж хүртэл. Сурагч мэдээлнэ, байршлаа сонгоно. Багш хүлээн авна, тусламжийг зохион байгуулна. Сурагч явцаа харна. Таван алхам, нэг урсгал. Энэ бол бодит дэлгэцүүд.",
  T(160, 100, 1600, 90, "Нэг дохионоос бодит тусламж хүртэл.", 60, NAVY, True),
  R(150, 290, 318, 650, "0B1322", radius=40),
  I(SHOT("review-after", "student-sos-390.png"), 164, 304, 290, 622, radius=30),
  T(150, 950, 320, 40, "Сурагч · Яаралтай SOS", 20, SKY_DEEP, True, align="c"),
  T(520, 560, 120, 80, "→", 72, SKY, True, align="c"),
  R(660, 280, 1110, 684, "0B1322", radius=18),
  I(SHOT("03-staff-dashboard-1440.png"), 672, 292, 1086, 660, radius=10),
  T(660, 970, 1110, 40, "Ээлжийн багш · Ирсэн хүсэлтүүд (яаралтай нь улаанаар эхэнд)", 20, SKY_DEEP, True, align="c"),
  T(160, 200, 1600, 50, "Сурагч мэдээлнэ  →  Байршил сонгоно  →  Багш хүлээн авна  →  Тусламж зохион байгуулна  →  Сурагч явцаа харна", 25, INK, False),
  T(160, 1020, 1600, 40, "See it. Signal safely. Get help.", 24, MUTED, False, align="c"),
)

# 05 — LIVE DEMO transition
S(CLOUD,
  "Одоо бодитоор үзүүлье. (ХӨТӨЧ РҮҮ ШИЛЖИНЭ — 80–90 секунд.) Тэмүүлэн апп-аа нээнэ. Байршлаа сонгоно эсвэл гарын үсэгтэй QR холбоосоор орно. Яаралтай SOS илгээнэ. Уянга багшийн самбарт ирнэ. Багш хүлээн авснаа батална. Хариуцагч томилно. Сурагч, багш мессеж солилцоно. Сурагч статусаа харна. (Бүх хэрэглэгч зохиомол.)",
  T(160, 110, 1000, 100, "Одоо бодитоор үзүүлье.", 72, NAVY, True),
  T(160, 230, 760, 50, "Амьд демо · бодит програм · зохиомол сургууль", 24, MUTED),
  T(160, 320, 380, 520, "1  Тэмүүлэн Happy2Help-ээ нээнэ\n\n2  Байршил сонгоно эсвэл гарын үсэгтэй QR холбоос\n\n3  Яаралтай SOS илгээнэ\n\n4  Багшийн самбарт хүсэлт ирнэ", 25, INK, False, line=1.25),
  T(540, 320, 360, 520, "5  Уянга багш хүлээн авснаа батална\n\n6  Хариуцагч ажилтан томилогдоно\n\n7  Сурагч, багш мессеж солилцоно\n\n8  Сурагч шинэчлэгдсэн статусаа харна", 25, INK, False, line=1.25),
  R(1000, 290, 790, 500, "0B1322", radius=18),
  I(SHOT("review-after", "teacher-inbox-1440.png"), 1010, 300, 770, 481, radius=10),
  R(960, 560, 262, 470, "0B1322", radius=34),
  I(SHOT("review-after", "student-home-390.png"), 971, 571, 240, 448, radius=26),
  R(160, 900, 700, 70, NAVY, radius=35),
  T(160, 900, 700, 70, "→  Хөтөч рүү шилжинэ · 80–90 сек", 26, WHITE, True, align="c", anchor="m"),
)

# 06 — MORE THAN AN SOS BUTTON (five real module screens)
mods = [
    ("review-after/student-chat-390.png", "Happy Connect", "сурагч, багшийн харилцаа"),
    ("review-after/student-care-390.png", "Happy Care", "эмч, нийгмийн ажилтанд хандах"),
    ("review-after/student-patrol-390.png", "Happy Patrol", "сурагчдын аюулгүй ээлж"),
    ("review-after/student-council-390.png", "Happy Council", "сурагчдын оролцоо, санал, шийдвэр"),
    ("08-student-ai-390.png", "Happy AI", "Gemma 4 · зөвхөн багшид санал"),
]
els6 = [T(160, 100, 1600, 90, "Тусламж бол зөвхөн нэг товч биш.", 60, NAVY, True),
        T(160, 195, 1600, 50, "Сургуулийн нийгэмлэгийг бүхэлд нь дэмжинэ.", 26, MUTED)]
for i, (p, name, sub) in enumerate(mods):
    x = 160 + i * 330
    if name == "Happy AI":
        els6 += [R(x - 10, 275, 290, 597, NAVY, radius=28),
                 T(x + 12, 300, 246, 40, "HAPPY AI · GEMMA 4", 16, SUN, True, spacing=2),
                 T(x + 12, 335, 246, 60, "Бодит дүгнэлт · 15:42", 22, WHITE, True),
                 T(x + 12, 380, 246, 60, "google/gemma-4-26b-a4b-it · 8.9 с · Zod ✓", 15, "C9D8EE"),
                 T(x + 12, 430, 246, 40, "Ангилал: ДЭЭРЭЛХЭЛТ · Яаралтай: ӨНДӨР", 15, MINT, True),
                 T(x + 12, 475, 246, 250, "«Долоо хоногийн турш ахлах ангийн хоёр хүү сурагчийг коридорт тулгаж, доромжилж, цүнх булааж байсан. Өнөөдөр мөн мөнгө нэхсэн тул сурагч айж байна.»", 16, WHITE, False, line=1.35),
                 T(x + 12, 735, 246, 60, "Зохиомол жишээ · зөвхөн багшид · шийдвэрийг хүн гаргана", 13, MUTED, False, line=1.3),
                 T(x + 12, 800, 246, 50, "docs/ai-proof/latest.json", 13, "C9D8EE", False, font="Consolas"),
                 T(x - 10, 895, 300, 44, name, 26, NAVY, True, align="c"),
                 T(x - 10, 940, 300, 70, sub, 19, INK, False, align="c")]
        continue
    els6 += [R(x - 10, 275, 290, 597, "0B1322", radius=36),
             I(SHOT(*p.split("/")), x + 2, 287, 266, 573, radius=28),
             T(x - 10, 895, 300, 44, name, 26, NAVY, True, align="c"),
             T(x - 10, 940, 300, 70, sub, 19, INK, False, align="c")]
els6.append(T(160, 1025, 1600, 36, "Бодит дэлгэцүүд · Happy AI карт нь бодит баталгаажсан гаралт (UI дэлгэц биш) · AI санал зөвхөн багшид очно, шийдвэрийг хүн гаргана", 18, MUTED, align="c"))
S(WHITE, "Тусламж бол зөвхөн нэг товч биш. Happy Connect: сурагч, багшийн харилцаа. Happy Care: эмч, нийгмийн ажилтан. Happy Patrol: сурагчдын аюулгүй ээлж. Happy Council: сурагчдын оролцоо, санал. Happy AI: Gemma 4 зөвхөн багшид санал өгнө, шийдвэрийг хүн гаргана.", *els6)

# 07 — WE ACTUALLY BUILT IT
S(NAVY_DEEP,
  "Зүгээр нэг санаа биш. 48-аас 48 end-to-end тест. 15-аас 15 амьд smoke шалгалт. SOS, хүлээн авалт, томилолт, хамгаалагдсан мессеж бүгд хадгалагдана. Сургуулиар хязгаарласан эрх, сөрөг тестүүд. Серверээр баталгаажсан гарын үсэгтэй QR. Gemma 4-ийн бодит дүгнэлт тусад нь баталгаажсан. Шударгаар: локал тест, production үйлчилгээ биш, erxes gateway холболт хүлээгдэж байна.",
  T(160, 90, 1600, 90, "Зүгээр нэг санаа биш.", 64, WHITE, True),
  T(160, 230, 420, 130, "48/48", 108, SUN, True),
  T(160, 370, 420, 40, "автомат E2E тест", 24, "C9D8EE"),
  T(600, 230, 420, 130, "15/15", 108, SUN, True),
  T(600, 370, 420, 40, "амьд web/BFF smoke шалгалт", 24, "C9D8EE"),
  T(160, 460, 880, 300, "✓  SOS, хүлээн авалт, томилолт, хамгаалагдсан мессеж — бүгд хадгалагдана\n✓  Сургуулиар хязгаарласан эрх · сөрөг (татгалзах) тестүүд\n✓  Серверээр баталгаажсан гарын үсэгтэй байршлын QR\n✓  Gemma 4 бодит дүгнэлт тусад нь баталгаажсан (15:42, Zod)", 25, WHITE, False, line=1.5),
  # mini architecture (right)
  R(1180, 240, 580, 90, "123F73", radius=18), T(1180, 240, 580, 90, "Сурагч / багшийн хөтөч", 26, WHITE, True, align="c", anchor="m"),
  T(1180, 330, 580, 40, "↓", 28, SKY, True, align="c"),
  R(1180, 370, 580, 90, "123F73", radius=18), T(1180, 370, 580, 90, "Next.js вэб + BFF (allow-list)", 26, WHITE, True, align="c", anchor="m"),
  T(1180, 460, 580, 40, "↓", 28, SKY, True, align="c"),
  R(1180, 500, 580, 110, "123F73", radius=18, line=SUN), T(1180, 500, 580, 110, "Happy2Help erxes API plugin\nFederation 2 · tenant/сургууль/үүрэг/кейс эрх", 24, WHITE, True, align="c", anchor="m"),
  T(1180, 610, 580, 40, "↓", 28, SKY, True, align="c"),
  R(1180, 650, 580, 90, "123F73", radius=18), T(1180, 650, 580, 90, "MongoDB · Redis бүртгэл · Gemma 4 (async)", 24, WHITE, True, align="c", anchor="m"),
  T(160, 840, 1600, 120, "Шударгаар: локал тестийн үр дүн · production яаралтай үйлчилгээ биш · erxes gateway / Core UI холболт хүлээгдэж байна · AI санал хүн шалгана, хязгаартай байж болно · бүх өгөгдөл зохиомол", 19, MUTED, False, line=1.4),
)

# 08 — THE ENDING
S(WHITE,
  "Эхэнд бид асуусан: тусламж хэрэгтэй үед хүүхэд хэнд хандах вэ? Хариулт нь: хэн нэгэн сонсоход бүх зүйл өөрчлөгдөж эхэлдэг. Хүүхэд бүр сонсогдох эрхтэй. Happy2Help. Баярлалаа.",
  T(160, 150, 1200, 180, "Хэн нэгэн сонсоход\nбүх зүйл өөрчлөгдөж эхэлдэг.", 54, INK, False, line=1.2),
  T(160, 420, 1200, 110, "Хүүхэд бүр сонсогдох эрхтэй.", 72, NAVY, True),
  I(BRAND["wordmark"], 160, 600, 560, 190, fit="contain"),
  I(BRAND["mascot"], 1330, 300, 440, 508, fit="contain"),
  T(160, 840, 900, 40, "Built with care by ZEVQORA.", 24, MUTED),
  T(160, 890, 900, 50, "See it. Signal safely. Get help.", 30, SKY_DEEP, True),
  T(160, 990, 1600, 40, "github.com/BeBecpp/Happy2Help_Hackathon · AGPL-3.0 · Hacktoberfest Hack Day Ulaanbaatar × erxes Hub 2026", 18, MUTED),
)

# ------------------------------------------------------------------ PPTX renderer
def px(v): return Emu(int(round(v * EMU_PER_PX)))
def rgb(h): return RGBColor.from_string(h)

def build_pptx(path):
    prs = Presentation()
    prs.slide_width = px(W); prs.slide_height = px(H)
    blank = prs.slide_layouts[6]
    for sd in slides:
        sl = prs.slides.add_slide(blank)
        bg = sl.background.fill; bg.solid(); bg.fore_color.rgb = rgb(sd["bg"])
        for e in sd["els"]:
            k = e["k"]
            if k == "rect":
                shp = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if e["radius"] else MSO_SHAPE.RECTANGLE, px(e["x"]), px(e["y"]), px(e["w"]), px(e["h"]))
                if e["radius"]:
                    shp.adjustments[0] = min(0.5, e["radius"] / min(e["w"], e["h"]))
                shp.fill.solid(); shp.fill.fore_color.rgb = rgb(e["fill"])
                if e["line"]:
                    shp.line.color.rgb = rgb(e["line"]); shp.line.width = Pt(2)
                else:
                    shp.line.fill.background()
                shp.shadow.inherit = False
            elif k in ("ring", "dot"):
                r = e["r"]
                shp = sl.shapes.add_shape(MSO_SHAPE.OVAL, px(e["cx"] - r), px(e["cy"] - r), px(2 * r), px(2 * r))
                if k == "dot":
                    shp.fill.solid(); shp.fill.fore_color.rgb = rgb(e["color"]); shp.line.fill.background()
                else:
                    shp.fill.background(); shp.line.color.rgb = rgb(e["color"]); shp.line.width = Pt(e["width"] * 0.75)
                shp.shadow.inherit = False
            elif k == "img":
                pic = sl.shapes.add_picture(e["path"], px(e["x"]), px(e["y"]), px(e["w"]), px(e["h"]))
                if e["fit"] == "cover":
                    # crop to the box aspect (keep top) — real pixels, no distortion
                    iw, ih = pic.image.size
                    box = e["w"] / e["h"]; img = iw / ih
                    if img > box:   # too wide → crop sides
                        keep = box / img; pic.crop_left = (1 - keep) / 2; pic.crop_right = (1 - keep) / 2
                    elif img < box: # too tall → crop bottom (keep top)
                        keep = img / box; pic.crop_bottom = 1 - keep
                else:
                    iw, ih = pic.image.size
                    box = e["w"] / e["h"]; img = iw / ih
                    if img > box:
                        nh = e["w"] / img; pic.height = px(nh); pic.top = px(e["y"] + (e["h"] - nh) / 2)
                    else:
                        nw = e["h"] * img; pic.width = px(nw); pic.left = px(e["x"] + (e["w"] - nw) / 2)
            elif k == "text":
                tb = sl.shapes.add_textbox(px(e["x"]), px(e["y"]), px(e["w"]), px(e["h"]))
                tf = tb.text_frame; tf.word_wrap = True
                tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
                tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[e["anchor"]]
                first = True
                for line in e["text"].split("\n"):
                    p = tf.paragraphs[0] if first else tf.add_paragraph(); first = False
                    p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[e["align"]]
                    p.line_spacing = e["line"]
                    run = p.add_run(); run.text = line
                    f = run.font; f.name = e["font"]; f.size = Pt(e["size"] * 0.75); f.bold = e["bold"]; f.color.rgb = rgb(e["color"])
                    if e["spacing"]:
                        rPr = run._r.get_or_add_rPr(); rPr.set("spc", str(int(e["spacing"] * 75)))
        sl.notes_slide.notes_text_frame.text = sd["notes"]
    prs.save(path)

# ------------------------------------------------------------------ HTML mirror (previews/PDF only)
def build_html(path):
    def css_text(e):
        al = {"l": "left", "c": "center", "r": "right"}[e["align"]]
        jc = {"t": "flex-start", "m": "center", "b": "flex-end"}[e["anchor"]]
        sp = f"letter-spacing:{e['spacing']}px;" if e["spacing"] else ""
        return (f"left:{e['x']}px;top:{e['y']}px;width:{e['w']}px;height:{e['h']}px;font-size:{e['size']}px;color:#{e['color']};"
                f"font-weight:{700 if e['bold'] else 400};text-align:{al};line-height:{e['line']};{sp}display:flex;flex-direction:column;justify-content:{jc};")
    pages = []
    for sd in slides:
        parts = []
        for e in sd["els"]:
            k = e["k"]
            if k == "rect":
                b = f"border:2px solid #{e['line']};" if e["line"] else ""
                parts.append(f"<div class=a style='left:{e['x']}px;top:{e['y']}px;width:{e['w']}px;height:{e['h']}px;background:#{e['fill']};border-radius:{e['radius']}px;{b}'></div>")
            elif k == "ring":
                r = e["r"]; parts.append(f"<div class=a style='left:{e['cx']-r}px;top:{e['cy']-r}px;width:{2*r}px;height:{2*r}px;border:{e['width']}px solid #{e['color']};border-radius:50%'></div>")
            elif k == "dot":
                r = e["r"]; parts.append(f"<div class=a style='left:{e['cx']-r}px;top:{e['cy']-r}px;width:{2*r}px;height:{2*r}px;background:#{e['color']};border-radius:50%'></div>")
            elif k == "img":
                src = "file:///" + e["path"].replace("\\", "/")
                fit = "cover" if e["fit"] == "cover" else "contain"
                parts.append(f"<div class=a style='left:{e['x']}px;top:{e['y']}px;width:{e['w']}px;height:{e['h']}px;border-radius:{e['radius']}px;overflow:hidden'><img src='{src}' style='width:100%;height:100%;object-fit:{fit};object-position:top center'></div>")
            elif k == "text":
                t = html.escape(e["text"]).replace("\n", "<br>")
                parts.append(f"<div class=a style=\"{css_text(e)}\"><div>{t}</div></div>")
        pages.append(f"<section class=slide style='background:#{sd['bg']}'>{''.join(parts)}</section>")
    doc = ("<!doctype html><meta charset='utf-8'><style>@page{size:1920px 1080px;margin:0}html,body{margin:0;padding:0;background:#222}"
           f".slide{{position:relative;width:{W}px;height:{H}px;overflow:hidden;font-family:'Segoe UI','Noto Sans',system-ui,sans-serif;page-break-after:always;break-after:page}}"
           ".a{position:absolute}</style>" + "".join(pages))
    with open(path, "w", encoding="utf-8") as f: f.write(doc)

if __name__ == "__main__":
    os.makedirs(OUT, exist_ok=True)
    for sd in slides:
        for e in sd["els"]:
            if e["k"] == "img" and not os.path.exists(e["path"]): sys.exit("missing image: " + e["path"])
    build_pptx(os.path.join(OUT, "Happy2Help-Final-Judges-Pitch.pptx"))
    build_html(os.path.join(OUT, "_mirror.html"))
    json.dump([{"slide": i + 1, "notes": s["notes"]} for i, s in enumerate(slides)], open(os.path.join(OUT, "_notes.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    print("built", len(slides), "slides ->", OUT)
